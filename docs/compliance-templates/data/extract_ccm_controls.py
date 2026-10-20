"""Turns the Cloud Security Alliance's Cloud Controls Matrix (CCM) into
the same import-ready CSV shape the FedRAMP/ISO 27001/SOC 2/C5 kits
use -- for CSA STAR certification directly, and as the base register
for Singapore government cloud work via MTCS SS 584 (see below).

## What to download yourself -- this repo ships none of it

Go to https://cloudsecurityalliance.org/artifacts/cloud-controls-matrix-v4
("Download this Resource"), sign in or create a free CSA account, and
download the CCM v4 bundle. Unzip it and find the CCM spreadsheet
itself (an .xlsx with one row per control, columns for Control Domain/
Control ID/Control Title/Control Specification plus several standards-
mapping columns you don't need here) -- save *that one file* as
ccm-v4.xlsx next to this script. CSA's own licensing FAQ
(https://cloudsecurityalliance.org/artifacts/ccm-licensing-faq) draws
the same line BSI's C5 license does: free for your own internal use,
licensed if you customize the CCM or embed it in a product. Running
this script against your own downloaded copy, for your own tenant, is
the former. This repo committing a derived CSV and pushing it to
GitHub on your behalf would be the latter -- which is exactly why,
like `extract_c5_controls.py` right below, only the *script* is
checked in here, never its output.

## A heads-up this script's siblings don't need

Every other extractor in this directory was built and run against a
real downloaded copy of its source before being committed. This one
wasn't -- CSA's download is account-gated, and no independently
verifiable copy of the current CCM v4 spreadsheet's exact column
headers was available while writing this. Column matching below is
deliberately loose (case-insensitive substring match against a few
likely header phrases, not exact strings) to absorb small wording
differences, but if your real file's headers don't match any
candidate, extraction will raise with the sheet's actual header row
printed -- add your file's real header text to the matching HEADERS_*
constants and re-run. Please open an issue (or just fix it) if CSA's
real headers turn out to need an adjustment this script doesn't cover.

## Singapore government cloud work (MTCS SS 584)

CSA STAR and Singapore's own MTCS SS 584 are different schemes --
MTCS, not CSA STAR by itself, is what Singapore government cloud
procurement actually requires (Level 3 for high-impact/government
systems) -- but IMDA formally cross-mapped the two in its own
CSA-STAR-to-MTCS cross-certification gap-analysis and implementation-
guideline documents (search "IMDA MTCS CSA STAR cross-certification"),
and MTCS SS 584 itself is a paid Singapore Standard (sold by Enterprise
Singapore) with no public control text at all -- unlike CCM, there's
nothing free this script could even extract. The practical path: run
this script for your CCM-based register, use it as the base for both
CSA STAR and the MTCS gap, and add IMDA's own cross-certification
documents' handful of MTCS-specific supplementary requirements by hand
on top -- same register, no second asset type needed.

Requires openpyxl (`pip install openpyxl`) -- already a RAIN backend
dependency, so it's on hand already if you're running this from a
checkout with the backend installed.

Usage: download and save the CCM spreadsheet as ccm-v4.xlsx next to
this script (see above), then run `python3 extract_ccm_controls.py`.
"""
from __future__ import annotations

import csv
import re

import openpyxl

IN_PATH = "ccm-v4.xlsx"
OUT_PATH = "ccm-v4-controls.csv"

# Loose, case-insensitive substring candidates for each column this
# script needs -- see the module docstring's "heads-up" section on why
# this isn't an exact-string match.
HEADERS_DOMAIN = ("control domain",)
HEADERS_ID = ("control id",)
HEADERS_TITLE = ("control title",)
HEADERS_SPEC = ("control specification", "control description")


def clean(text: str | None) -> str:
    return re.sub(r"\s+", " ", (text or "")).strip()


def _find_header_row_and_columns(ws) -> tuple[int, dict[str, int]]:
    """Scans the first 20 rows for one that looks like a header row (at
    least Control ID and Control Title/Specification present), rather
    than assuming row 1 -- CCM spreadsheets in the wild often carry a
    title/version banner above the real header row."""
    for row_idx in range(1, 21):
        cells = [clean(str(c.value)) if c.value is not None else "" for c in ws[row_idx]]
        lower = [c.lower() for c in cells]

        def find(candidates: tuple[str, ...]) -> int | None:
            for i, cell in enumerate(lower):
                if any(cand in cell for cand in candidates):
                    return i
            return None

        col_id = find(HEADERS_ID)
        col_title = find(HEADERS_TITLE)
        col_spec = find(HEADERS_SPEC)
        col_domain = find(HEADERS_DOMAIN)
        if col_id is not None and (col_title is not None or col_spec is not None):
            columns = {"id": col_id, "domain": col_domain, "title": col_title, "spec": col_spec}
            return row_idx, columns

    header_preview = [clean(str(c.value)) if c.value is not None else "" for c in ws[1]]
    raise SystemExit(
        "Couldn't find a header row with recognizable Control ID/Title/Specification columns "
        f"in the first 20 rows of sheet {ws.title!r}. Row 1 reads: {header_preview!r} -- "
        "open the file, find your real header row, and add its exact column phrases to this "
        "script's HEADERS_* constants."
    )


def _find_sheet(wb) -> "openpyxl.worksheet.worksheet.Worksheet":
    for name in wb.sheetnames:
        if "ccm" in name.lower():
            return wb[name]
    # Fall back to the first sheet if nothing's named obviously --
    # some CCM releases ship the controls on an unlabeled or
    # differently-named first tab.
    return wb[wb.sheetnames[0]]


def extract(xlsx_path: str) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = _find_sheet(wb)
    header_row, columns = _find_header_row_and_columns(ws)

    rows = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):

        def cell(key: str) -> str:
            idx = columns.get(key)
            return clean(str(row[idx])) if idx is not None and idx < len(row) and row[idx] is not None else ""

        control_id = cell("id")
        if not control_id:
            continue  # blank row, or a domain-header/section-break row with no control of its own
        title = cell("title")
        spec = cell("spec")
        rows.append({
            "Name": f"{control_id} {title}".strip(),
            "External ID": control_id,
            "Control ID": control_id,
            "Statement ID": "",
            "Control Question": spec or title,
        })
    return rows


if __name__ == "__main__":
    rows = extract(IN_PATH)
    with open(OUT_PATH, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["Name", "External ID", "Control ID", "Statement ID", "Control Question"])
        w.writeheader()
        w.writerows(rows)
    print(f"wrote {OUT_PATH}: {len(rows)} rows")
